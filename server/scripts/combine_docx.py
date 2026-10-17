"""Cover letter + résumé in one Word file: the cover letter's content, a page
break, then the résumé exactly as built.

    python combine_docx.py <cover_letter.docx> <resume.docx> <out.docx>

The résumé is the base document, so its styles, bullets (numbering), margins
and header carry over untouched; the cover letter's paragraphs and tables are
copied in front of it. Styles the résumé lacks are copied along; images and
links are re-attached to the new file. Anything else the cover letter links
to (charts, embedded objects) can't be carried safely, so the script stops
with exit code 2 and writes nothing rather than a file Word would call
corrupted.
"""
import copy
import io
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


class Unsupported(Exception):
    pass


def reattach_relationships(element, src_part, dst_part):
    """Point r:id / r:embed / r:link attributes at relationships in dst_part."""
    for node in element.iter():
        for attr, rid in list(node.attrib.items()):
            if not attr.startswith(R_NS):
                continue
            rel = src_part.rels.get(rid)
            if rel is None:
                continue
            if rel.is_external:
                node.set(attr, dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True))
            elif rel.reltype == RT.IMAGE:
                # Added as a new image part of the résumé's package, so two
                # "image1.png" files can't collide.
                new_rid, _ = dst_part.get_or_add_image(io.BytesIO(rel.target_part.blob))
                node.set(attr, new_rid)
            else:
                raise Unsupported(f"cover letter links to a {rel.reltype.rsplit('/', 1)[-1]} part")


def combine(cover_path, resume_path, out_path):
    cover = Document(cover_path)
    resume = Document(resume_path)

    known = {s.get(qn("w:styleId")) for s in resume.styles.element.findall(qn("w:style"))}
    for style in cover.styles.element.findall(qn("w:style")):
        if style.get(qn("w:styleId")) not in known:
            resume.styles.element.append(copy.deepcopy(style))

    body = resume.element.body
    anchor = body[0]
    for element in cover.element.body:
        if element.tag == qn("w:sectPr"):
            continue  # the résumé's page setup applies to the whole file
        clone = copy.deepcopy(element)
        # A cover-letter list would point at the résumé's numbering
        # definitions; the text stays, the list formatting doesn't.
        for num_pr in list(clone.iter(qn("w:numPr"))):
            num_pr.getparent().remove(num_pr)
        reattach_relationships(clone, cover.part, resume.part)
        anchor.addprevious(clone)

    page_break = resume.add_paragraph()  # appended at the end, then moved
    page_break.add_run().add_break(WD_BREAK.PAGE)
    anchor.addprevious(page_break._p)

    resume.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print("usage: combine_docx.py <cover_letter.docx> <resume.docx> <out.docx>", file=sys.stderr)
        sys.exit(64)
    try:
        combine(*sys.argv[1:])
    except Unsupported as err:
        print(f"not combined: {err}", file=sys.stderr)
        sys.exit(2)
